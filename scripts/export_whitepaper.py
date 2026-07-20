#!/usr/bin/env python
"""Export the whitepaper and manifesto markdown files to docx AND pdf, each
with a Le Kiosque cover page.

The cover reproduces the daily/weekly PDF reports (report_pdf.py): green
masthead band with the platform name and tagline, document title over an
accent rule, edition date, author and generation lines.

docx: pandoc + python-docx (cover injected as styled paragraphs).
pdf : pandoc → styled HTML with an HTML/CSS cover → headless Chrome print.

    .venv/bin/python scripts/export_whitepaper.py
"""

import subprocess
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"

# Report palette (report_pdf.py / lekiosque_theme.py)
BAND, INK, ACCENT, MUTED = "0E5432", "1C241F", "166F43", "6E7A72"
FONT = "Helvetica"

TEXTS = {
    "fr": {
        "md": "whitepaper-fr.md",
        "docx": "whitepaper-fr.docx",
        "tagline": "Plateforme d'intelligence documentaire de la presse gabonaise",
        "title": "Livre blanc",
        "subtitle": "Une plateforme souveraine d'intelligence documentaire pour la presse gabonaise",
        "edition": lambda now: "Version 2.1 · Juillet 2026",
        "author": "Auteur : Vhiny Mombo",
        "generated": lambda now: f"Document généré le {now.strftime('%d/%m/%Y')}",
        "drop_headings": ("Le Kiosque", "Une plateforme souveraine",
                          "Livre blanc ·", "Auteur :"),
    },
    "en": {
        "md": "whitepaper-en.md",
        "docx": "whitepaper-en.docx",
        "tagline": "Document-intelligence platform for the Gabonese press",
        "title": "White paper",
        "subtitle": "A sovereign document-intelligence platform for the Gabonese press",
        "edition": lambda now: "Version 2.1 · July 2026",
        "author": "Author: Vhiny Mombo",
        "generated": lambda now: f"Document generated on {now.strftime('%B %d, %Y')}",
        "drop_headings": ("Le Kiosque", "A sovereign document-intelligence",
                          "White paper ·", "Author:"),
    },
    "manifeste-fr": {
        "md": "manifeste-fr.md",
        "docx": "manifeste-fr.docx",
        "tagline": "Plateforme d'intelligence documentaire de la presse gabonaise",
        "title": "Manifeste",
        "subtitle": "Une plateforme souveraine d'intelligence documentaire pour la presse gabonaise",
        "edition": lambda now: "Version 1.0 · Juillet 2026",
        "author": "Auteur : Vhiny Mombo",
        "generated": lambda now: f"Document généré le {now.strftime('%d/%m/%Y')}",
        "drop_headings": ("Le Kiosque", "Une plateforme souveraine",
                          "Manifeste ·", "Auteur :"),
    },
    "manifesto-en": {
        "md": "manifesto-en.md",
        "docx": "manifesto-en.docx",
        "tagline": "Document-intelligence platform for the Gabonese press",
        "title": "Manifesto",
        "subtitle": "A sovereign document-intelligence platform for the Gabonese press",
        "edition": lambda now: "Version 1.0 · July 2026",
        "author": "Author: Vhiny Mombo",
        "generated": lambda now: f"Document generated on {now.strftime('%B %d, %Y')}",
        "drop_headings": ("Le Kiosque", "A sovereign document-intelligence",
                          "Manifesto ·", "Author:"),
    },
}


def _style_run(run, size, bold, color):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _shade(paragraph, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shd)


def _bottom_border(paragraph, color, size_eighths=24):
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size_eighths))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    paragraph._p.get_or_add_pPr().append(pbdr)


def build_cover(doc, t, now):
    """Append cover paragraphs to doc and return them (moved to the top later)."""
    paras = []

    def para(text="", size=11, bold=False, color=INK, shade=None,
             before=0, after=0):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        if shade:
            _shade(p, shade)
        if text:
            _style_run(p.add_run(text), size, bold, color)
        paras.append(p)
        return p

    # Masthead band, like the PDF's green header block
    para(" ", size=8, shade=BAND)
    para("Le Kiosque", size=34, bold=True, color="FFFFFF", shade=BAND, after=6)
    para(t["tagline"], size=13, color="FFFFFF", shade=BAND)
    para(" ", size=8, shade=BAND)

    # Push the title block toward the middle of the page
    para(after=170)

    para(t["title"], size=27, bold=True, color=INK, after=10)
    rule = para(after=26)
    rule.paragraph_format.right_indent = Cm(10.5)
    _bottom_border(rule, ACCENT)
    para(t["subtitle"], size=15, color=INK, after=8)
    para(t["edition"](now), size=11.5, color=MUTED)

    # Bottom block: author + generation notes
    para(after=110)
    para(t["author"], size=12, bold=True, color=INK, after=4)
    para(t["generated"](now), size=10, color=MUTED)

    # Page break onto the document body
    br = doc.add_paragraph()
    br.add_run().add_break(WD_BREAK.PAGE)
    paras.append(br)
    return paras


def export(lang):
    t = TEXTS[lang]
    md, out = DOCS / t["md"], DOCS / t["docx"]
    # cwd=DOCS so relative image paths in the markdown resolve
    subprocess.run(["pandoc", t["md"], "-o", t["docx"]], check=True, cwd=DOCS)

    doc = Document(str(out))

    # The cover carries the title, subtitle and version line; drop them from
    # the body along with the separator rule that followed them
    for p in list(doc.paragraphs[:4]):
        if p.text.strip().startswith(t["drop_headings"]) or not p.text.strip():
            p._p.getparent().remove(p._p)

    anchor = doc.element.body[0]
    for p in build_cover(doc, t, datetime.now()):
        anchor.addprevious(p._p)

    doc.save(str(out))
    print(f"✅ {out.name}")


CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

PRINT_CSS = """
@page { size: A4; margin: 20mm 18mm; }
* { -webkit-print-color-adjust: exact; print-color-adjust: exact; }
body { font-family: Helvetica, Arial, sans-serif; color: #1c241f;
       font-size: 10.5pt; line-height: 1.55; margin: 0; }
h2 { color: #166f43; font-size: 16pt; margin: 1.4em 0 0.5em; line-height: 1.25; }
h3 { font-size: 12.5pt; margin: 1.2em 0 0.4em; }
a { color: #166f43; }
img { max-width: 100%; }
table { border-collapse: collapse; width: 100%; font-size: 9.5pt; margin: 0.8em 0; }
th, td { border: 1px solid #d8ded9; padding: 5px 8px; text-align: left; }
th { background: #eef4f0; }
blockquote { font-style: italic; color: #6e7a72; border-left: 3px solid #166f43;
             margin: 1em 0; padding-left: 12px; }
pre { background: #f4f1e9; padding: 10px 12px; border-radius: 4px; }
code { font-family: Menlo, monospace; font-size: 9pt; }
hr { border: none; border-top: 1px solid #d8ded9; margin: 1.6em 0; }

.cover { height: 257mm; display: flex; flex-direction: column;
         page-break-after: always; }
.band { background: #0e5432; color: #ffffff; padding: 14mm 12mm 10mm; }
.band h1 { color: #ffffff; font-size: 26pt; margin: 0 0 4mm; }
.band p { margin: 0; font-size: 11.5pt; }
.cover-mid { margin-top: 52mm; }
.cover-mid h2 { color: #1c241f; font-size: 22pt; margin: 0 0 4mm; }
.rule { width: 45mm; border-bottom: 2.5pt solid #166f43; margin-bottom: 8mm; }
.sub { font-size: 12.5pt; margin: 0 0 3mm; }
.ed { color: #6e7a72; font-size: 10.5pt; margin: 0; }
.cover-foot { margin-top: auto; }
.author { font-weight: bold; margin: 0 0 2mm; }
.gen { color: #6e7a72; font-size: 9.5pt; margin: 0; }
"""


def export_pdf(key):
    t = TEXTS[key]
    now = datetime.now()
    body = subprocess.run(
        ["pandoc", t["md"], "-t", "html"],
        capture_output=True, text=True, check=True, cwd=DOCS,
    ).stdout
    # The HTML cover replaces the title block (everything up to the first rule)
    body = body.split("<hr />", 1)[-1]

    cover = f"""
<div class="cover">
  <div class="band"><h1>Le Kiosque</h1><p>{t["tagline"]}</p></div>
  <div class="cover-mid">
    <h2>{t["title"]}</h2>
    <div class="rule"></div>
    <p class="sub">{t["subtitle"]}</p>
    <p class="ed">{t["edition"](now)}</p>
  </div>
  <div class="cover-foot">
    <p class="author">{t["author"]}</p>
    <p class="gen">{t["generated"](now)}</p>
  </div>
</div>"""

    html = (f"<!doctype html><html><head><meta charset='utf-8'>"
            f"<style>{PRINT_CSS}</style></head><body>{cover}{body}</body></html>")
    tmp = DOCS / f".print-{key}.html"
    tmp.write_text(html, encoding="utf-8")
    pdf = DOCS / t["docx"].replace(".docx", ".pdf")
    try:
        subprocess.run(
            [CHROME, "--headless", "--disable-gpu", "--no-pdf-header-footer",
             f"--print-to-pdf={pdf}", tmp.as_uri()],
            check=True, capture_output=True,
        )
    finally:
        tmp.unlink(missing_ok=True)
    print(f"✅ {pdf.name}")


if __name__ == "__main__":
    for key in TEXTS:
        export(key)
        export_pdf(key)
